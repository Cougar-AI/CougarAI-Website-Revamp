from __future__ import annotations

import json
import math
import os
import subprocess
import sys
from pathlib import Path


class ExtractionFailed(Exception):
    pass


def _rgb_hidden(color: int) -> bool:
    r, g, b = (color >> 16) & 255, (color >> 8) & 255, color & 255
    return min(r, g, b) >= 250  # white / near-white


def extract_pdf(path: str) -> dict:
    import pymupdf

    doc = pymupdf.open(path)
    parts: list[str] = []
    hidden = images = links = left_cols = right_cols = 0
    has_tables = False
    for page in doc:
        W, H = page.rect.width, page.rect.height
        # Unclipped, so text placed outside the page is seen (and counted as hidden) instead of silently dropped.
        for block in page.get_text("dict", clip=pymupdf.INFINITE_RECT())["blocks"]:
            if block.get("type") != 0:
                continue
            lines_out = []
            for line in block["lines"]:
                visible = []
                for sp in line["spans"]:
                    x0, y0, x1, y1 = sp["bbox"]
                    offpage = x1 < 0 or y1 < 0 or x0 > W or y0 > H
                    if sp["text"].strip() and (_rgb_hidden(sp["color"]) or sp["size"] < 3 or offpage):
                        hidden += len(sp["text"].strip())
                        continue
                    visible.append(sp["text"])
                t = "".join(visible).strip()
                if t:
                    lines_out.append(t)
            if lines_out:
                block_text = "\n".join(lines_out)
                parts.append(block_text)
                if len(block_text) >= 60:  # ignore short right-aligned date blocks
                    bx0, _, bx1, _ = block["bbox"]
                    if bx1 < 0.55 * W:
                        left_cols += 1
                    elif bx0 > 0.45 * W:
                        right_cols += 1
        images += len(page.get_images())
        links += len(page.get_links())
        try:
            has_tables = has_tables or len(page.find_tables().tables) > 0
        except Exception:
            pass
    pages = doc.page_count
    doc.close()
    return {
        "text": "\n".join(parts), "pages": pages, "hidden_removed": hidden, "images": images,
        "links": links, "has_tables": has_tables, "has_textboxes": False,
        "multi_column": left_cols >= 2 and right_cols >= 2, "header_footer_chars": 0,
    }


def _run_hidden(run) -> bool:
    f = run.font
    if f.hidden:
        return True
    if f.size is not None and f.size.pt < 3:
        return True
    try:
        rgb = f.color.rgb if f.color is not None and f.color.type is not None else None
    except Exception:
        rgb = None
    return rgb is not None and min(rgb) >= 250


def extract_docx(path: str) -> dict:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(path)
    hidden = 0

    def para_text(p: Paragraph) -> str:
        nonlocal hidden
        out = []
        for r in p.runs:
            if r.text.strip() and _run_hidden(r):
                hidden += len(r.text.strip())
                continue
            out.append(r.text)
        return "".join(out).strip()

    lines: list[str] = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = para_text(Paragraph(child, d))
            if t:
                lines.append(t)
        elif tag == "tbl":
            seen = set()
            for row in Table(child, d).rows:
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    for p in cell.paragraphs:
                        t = para_text(p)
                        if t:
                            lines.append(t)

    hf = 0
    for s in d.sections:
        for p in list(s.header.paragraphs) + list(s.footer.paragraphs):
            hf += len(p.text.strip())

    text = "\n".join(lines)
    words = len(text.split())
    return {
        "text": text, "pages": max(1, math.ceil(words / 500)),  # DOCX has no fixed pages; estimate
        "hidden_removed": hidden, "images": len(d.inline_shapes), "links": 0,
        "has_tables": len(d.tables) > 0,
        "has_textboxes": len(d.element.body.xpath(".//w:txbxContent")) > 0,
        "multi_column": False, "header_footer_chars": hf,
    }


def extract(path: str, kind: str) -> dict:
    return extract_pdf(path) if kind == "pdf" else extract_docx(path)


def extract_sandboxed(path: str | Path, kind: str, timeout: int = 25) -> dict:
    root = Path(__file__).resolve().parent.parent

    def _limits():  # POSIX only
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (1536 * 1024 * 1024,) * 2)  # memory
        resource.setrlimit(resource.RLIMIT_CPU, (20, 20))                  # CPU seconds

    try:
        proc = subprocess.run(
            [sys.executable, "-m", "ats.extract", str(path), kind],
            capture_output=True, timeout=timeout, cwd=root,
            env={"PATH": os.environ.get("PATH", ""), "PYTHONPATH": str(root)},  # no inherited secrets
            preexec_fn=_limits if os.name == "posix" else None,
        )
    except subprocess.TimeoutExpired:
        raise ExtractionFailed("parser timed out")
    if proc.returncode != 0 or len(proc.stdout) > 3_000_000:
        raise ExtractionFailed("parser failed")
    try:
        # Parse only the LAST line: libraries sometimes print warnings to stdout.
        return json.loads(proc.stdout.decode("utf-8", "ignore").strip().splitlines()[-1])
    except (json.JSONDecodeError, IndexError):
        raise ExtractionFailed("parser returned invalid output")


if __name__ == "__main__":  # worker entry point
    result = extract(sys.argv[1], sys.argv[2])
    sys.stdout.write("\n" + json.dumps(result) + "\n")

